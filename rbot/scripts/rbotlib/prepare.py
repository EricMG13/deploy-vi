"""`rbot prepare`: page-anchored text for every inventoried file (interfaces spec section 8.1).

Each readable file becomes `prepared/<sha16>-<slug>.txt`, its pages opened by
`<<< page N >>>` lines, so a citation's page locator and excerpt can be checked.
`prepared/index.json` records per file the prepared path, page count, a 400-character
first-page excerpt, the entity names and dates found on the first page, sheet names
and the `image_read` flag. A file that cannot be read is listed with `unreadable`
and the reason; the others are still prepared (exit 0).
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PyPdfError

from rbotlib import portfolio
from rbotlib.inventory import RunError
from rbotlib.rules import Rules, load_rules

SCHEMA = "rbot-prepared/1"
EXCERPT_CHARS = 400
MAX_ENTITIES = 8
MIN_PDF_TEXT = 20  # fewer extracted characters in the whole PDF: a scan, read as an image


class UnreadableError(Exception):
    """This file cannot be prepared; the message is recorded as `unreadable`."""


Sheet = tuple[str, list[list[str]]]  # sheet name, rows of cleaned cells


@dataclass
class Prepared:
    pages: list[str]
    sheets: list[Sheet] = field(default_factory=list)
    image_read: bool = False


def page_marker(n: int) -> str:
    return f"<<< page {n} >>>"


# --- readers -----------------------------------------------------------------


def read_text(path: Path) -> Prepared:
    """Plain text and Markdown: form feeds separate pages; without them, one page."""
    text = path.read_text(encoding="utf-8", errors="replace").replace("\r\n", "\n")
    return Prepared(text.split("\f"))


class _HTMLText(HTMLParser):
    BLOCKS = frozenset({"p", "div", "br", "tr", "li", "h1", "h2", "h3", "h4", "table", "section"})

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in ("script", "style"):
            self.hidden += 1
        elif tag in self.BLOCKS:
            self.parts.append("\n")
        elif tag in ("td", "th"):
            self.parts.append(" | ")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style"):
            self.hidden -= 1

    def handle_data(self, data: str) -> None:
        if not self.hidden:
            self.parts.append(data)


def read_html(path: Path) -> Prepared:
    parser = _HTMLText()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    lines = (" ".join(line.split()) for line in "".join(parser.parts).splitlines())
    return Prepared(["\n".join(line for line in lines if line)])


def read_pdf(path: Path) -> Prepared:
    """PDF pages through pypdf; a PDF with (almost) no extractable text is a scan."""
    try:
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise UnreadableError("encrypted PDF: a password is needed")
        pages = [(page.extract_text() or "").replace("\r\n", "\n") for page in reader.pages]
    except (PyPdfError, ValueError, KeyError, OSError) as exc:
        raise UnreadableError(f"pypdf could not read the PDF: {type(exc).__name__}: {exc}") from exc
    if sum(len(p.strip()) for p in pages) < MIN_PDF_TEXT:
        return Prepared(["" for _ in pages], image_read=True)
    return Prepared(pages)


_ERRORS = frozenset({"#VALUE!", "#DIV/0!", "#NAME?", "#REF!", "#NUM!", "#NULL!", "—"})


def clean_cell(value: object) -> str:
    """One cell as text: spreadsheet errors, `#N/A…`, `—` and blanks are an empty cell."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return (
            value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
        )
    text = str(value).strip()
    return "" if text in _ERRORS or text.startswith("#N/A") else text


def _trimmed(rows: list[list[str]]) -> list[list[str]]:
    out = []
    for row in rows:
        while row and not row[-1]:
            row = row[:-1]
        out.append(row)
    while out and not out[-1]:
        out.pop()
    return out


def _sheet_pages(sheets: list[Sheet]) -> Prepared:
    pages = []
    for name, rows in sheets:
        buffer = io.StringIO()
        csv.writer(buffer, lineterminator="\n").writerows(rows)
        pages.append(f"<<< sheet {name} >>>\n{buffer.getvalue()}")
    return Prepared(pages, sheets)


def read_csv(path: Path) -> Prepared:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    rows = [[clean_cell(c) for c in row] for row in csv.reader(io.StringIO(text))]
    return _sheet_pages([(path.stem, _trimmed(rows))])


def read_xlsx(path: Path) -> Prepared:
    """Every sheet's stored values (formula results as cached, never recalculated)."""
    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, ValueError, OSError) as exc:
        raise UnreadableError(f"openpyxl could not read the workbook: {exc}") from exc
    try:
        sheets = [
            (
                ws.title,
                _trimmed([[clean_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]),
            )
            for ws in book.worksheets
        ]
    finally:
        book.close()
    return _sheet_pages(sheets)


def read_image(path: Path) -> Prepared:
    """An image has no text: one empty page, flagged so its figures are Image-Read Values."""
    return Prepared([""], image_read=True)


READERS: dict[str, Callable[[Path], Prepared]] = {
    "txt": read_text,
    "md": read_text,
    "html": read_html,
    "pdf": read_pdf,
    "csv": read_csv,
    "xlsx": read_xlsx,
    "image": read_image,
}


def read(path: Path, kind: str) -> Prepared:
    reader = READERS.get(kind)
    if reader is None:
        raise UnreadableError(f"no reader for .{path.suffix.lower().lstrip('.')} files")
    return reader(path)


# --- first-page detection ------------------------------------------------------

_REGISTRANT = re.compile(r"^\s*(.+?)\s*\n\s*\(exact name of registrant", re.I | re.M)
_SUFFIX = (
    r"(?:Inc\.?|INC\.?|Corporation|CORPORATION|Corp\.|CORP\.|Company|COMPANY|plc|PLC"
    r"|Limited|LIMITED|Ltd\.?|LTD\.?|LLC|L\.P\.|N\.A\.)"
)
_NAME = re.compile(r"\b((?:[A-Z][\w&'\u2019.-]*,?\s+){1,6}" + _SUFFIX + r")(?![\w])")
_GENERIC = frozenset({"the", "this", "our", "such", "each", "any"})
_MONTHS = "january february march april may june july august september october november december"
_MONTH = {m: n for n, m in enumerate(_MONTHS.split(), 1)}
_MDY = re.compile(r"\b(" + "|".join(_MONTH) + r")\s+(\d{1,2}),\s*(\d{4})\b", re.I)
_DMY = re.compile(r"\b(\d{1,2})\s+(" + "|".join(_MONTH) + r")\s+(\d{4})\b", re.I)
_ISO = re.compile(r"\b(\d{4})-(\d{2})-(\d{2})(?!\d)")


def _clean(raw: str) -> str | None:
    """Trim a candidate name to the company: drop a preceding sentence and lower-case lead-ins."""
    words = raw.replace("|", " ").split()
    breaks = [i for i, w in enumerate(words[:-1]) if w.endswith(".") and len(w) > 4]
    words = words[breaks[-1] + 1 :] if breaks else words
    if any(w.isupper() and len(w) > 1 for w in words[:-1]):
        while words and not words[0].isupper():
            words.pop(0)
    if len(words) < 2 or {w.casefold() for w in words[:-1]} <= _GENERIC:
        return None
    return " ".join(words).strip(" ,")


def entities(first_page: str) -> list[str]:
    """Entity names on the first page: the SEC registrant line first, then corporate names."""
    found = [m.group(1) for m in _REGISTRANT.finditer(first_page)]
    found += [m.group(1) for m in _NAME.finditer(first_page)]
    names: list[str] = []
    for raw in found:
        name = _clean(raw)
        if name and name.casefold() not in {n.casefold() for n in names}:
            names.append(name)
    return names[:MAX_ENTITIES]


def _iso(y: str, m: int, d: str) -> str | None:
    try:
        return date(int(y), m, int(d)).isoformat()
    except ValueError:
        return None


def dates(first_page: str) -> list[str]:
    """ISO dates written on the first page, in the order they appear, without repeats."""
    hits = [(m.start(), _iso(m[3], _MONTH[m[1].lower()], m[2])) for m in _MDY.finditer(first_page)]
    hits += [(m.start(), _iso(m[3], _MONTH[m[2].lower()], m[1])) for m in _DMY.finditer(first_page)]
    hits += [(m.start(), _iso(m[1], int(m[2]), m[3])) for m in _ISO.finditer(first_page)]
    ordered = [d for _, d in sorted(hits) if d]
    return list(dict.fromkeys(ordered))


# --- the command ---------------------------------------------------------------


def slug(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:60]


def prepared_name(entry: dict[str, Any]) -> str:
    return f"prepared/{entry['sha256'][:16]}-{slug(Path(entry['path']).name)}.txt"


def anchored_text(pages: list[str]) -> str:
    """Each page under its `<<< page N >>>` line; a page without text is the line alone."""
    bodies = (text.rstrip() for text in pages)
    return "".join(f"{page_marker(n)}\n" + (f"{b}\n" if b else "") for n, b in enumerate(bodies, 1))


def prepare_file(run: Path, pack: Path, entry: dict[str, Any]) -> dict[str, Any]:
    record: dict[str, Any] = {"path": entry["path"], "sha256": entry["sha256"]}
    record |= {"kind": entry["kind"], "prepared": None, "pages": 0, "excerpt": ""}
    record |= {"entities": [], "dates": [], "sheets": [], "image_read": False, "unreadable": None}
    try:
        result = read(pack / entry["path"], entry["kind"])
    except UnreadableError as exc:
        record["unreadable"] = str(exc)
        return record
    name = prepared_name(entry)
    (run / name).write_text(anchored_text(result.pages), encoding="utf-8")
    first = result.pages[0] if result.pages else ""
    sheets = [n for n, _ in result.sheets]
    record |= {"prepared": name, "pages": len(result.pages), "sheets": sheets}
    record |= {"excerpt": " ".join(first.split())[:EXCERPT_CHARS], "image_read": result.image_read}
    record |= {"entities": entities(first), "dates": dates(first)}
    return record


def _portfolios(run: Path, pack: Path, files: list[dict[str, Any]], rules: Rules) -> dict[str, Any]:
    """Every portfolio found in the pack's sheets, written to prepared/portfolios/."""
    found: dict[str, portfolio.Portfolio] = {}
    notes = []
    for entry in (e for e in files if e["kind"] in ("csv", "xlsx")):
        try:
            sheets = read(pack / entry["path"], entry["kind"]).sheets
        except UnreadableError:
            continue  # already listed as unreadable in index.json
        for name, rows in sheets:
            label = entry["path"] if entry["kind"] == "csv" else f"{entry['path']} sheet {name}"
            note = portfolio.read_sheet(label, rows, found)
            notes += [note] if note else []
    return {"portfolios": portfolio.write_all(run, found, rules), "portfolio_notes": notes}


def write(run: Path, rules: Rules, only: str | None = None) -> dict[str, Any]:
    try:
        inventory = json.loads((run / "inventory.json").read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise RunError(f"{run}: no readable inventory.json; run rbot inventory first") from exc
    entries = [e for e in inventory["files"] if only is None or e["path"] == only]
    if only is not None and not entries:
        raise RunError(f"{only}: not in inventory.json")
    (run / "prepared").mkdir(exist_ok=True)
    index_path = run / "prepared" / "index.json"
    kept = {}
    if only is not None and index_path.is_file():
        kept = {r["path"]: r for r in json.loads(index_path.read_text(encoding="utf-8"))["files"]}
    pack = Path(inventory["pack_root"])
    for entry in entries:
        kept[entry["path"]] = prepare_file(run, pack, entry)
    index: dict[str, Any] = {"schema": SCHEMA}
    index["generated_at"] = datetime.now(UTC).isoformat(timespec="seconds")
    index["files"] = [kept[p] for p in sorted(kept)]
    index |= _portfolios(run, pack, inventory["files"], rules)
    index_path.write_text(json.dumps(index, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return index


def load_index(run: Path) -> dict[str, Any]:
    data: dict[str, Any] = json.loads((run / "prepared" / "index.json").read_text(encoding="utf-8"))
    return data


def run_command(args: argparse.Namespace) -> int:
    try:
        index = write(Path(args.run_dir), load_rules(args.root), args.only)
    except (RunError, OSError) as exc:
        print(exc)
        return 1
    unreadable = [r for r in index["files"] if r["unreadable"]]
    for r in unreadable:
        print(f"unreadable: {r['path']}: {r['unreadable']}")
    for note in index["portfolio_notes"]:
        print(f"portfolio: {note}")
    found = ", ".join(p["portfolio_id"] for p in index["portfolios"]) or "none"
    print(f"prepared/index.json: {len(index['files'])} file(s), {len(unreadable)} unreadable")
    print(f"portfolios: {found}")
    return 0


def add_command(sub: argparse._SubParsersAction[argparse.ArgumentParser]) -> None:
    c = sub.add_parser("prepare", help="write page-anchored text for every inventoried file")
    c.add_argument("--run", dest="run_dir", required=True, help="the run folder")
    c.add_argument("--only", help="prepare one file (its path as in inventory.json)")
    c.set_defaults(run=run_command)
