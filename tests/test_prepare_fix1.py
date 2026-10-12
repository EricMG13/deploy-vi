"""02G fix round 1: display scaling, per-file aborts, line breaks, index safety, PDF warnings."""

from __future__ import annotations

import io
import json
import re
import zipfile
from collections.abc import Callable
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from hypothesis import given
from hypothesis import strategies as st
from openpyxl import Workbook

from builders import FONT, make_pdf, make_run
from rbotlib.cli import main
from rbotlib.inventory import write_atomic
from rbotlib.prepare import anchored_text, cell_text, clean_cell, load_index
from rbotlib.validate import pages


def _prep(tmp_path: Path, files: dict[str, bytes], *extra: str) -> tuple[Path, dict[str, Any], int]:
    pack = tmp_path / "pack"
    pack.mkdir(exist_ok=True)
    for name, data in files.items():
        (pack / name).write_bytes(data)
    run = (
        make_run(pack)
        if not (pack / "rbot-runs").exists()
        else next((pack / "rbot-runs").iterdir())
    )
    assert main(["inventory", "--run", str(run)]) == 0
    code = main(["prepare", "--run", str(run), *extra])
    index = load_index(run) if (run / "prepared/index.json").exists() else {"files": []}
    if (run / "prepared/index.json").exists():
        assert index["schema"] == "rbot-prepared/1"
        assert datetime.fromisoformat(index["generated_at"]).tzinfo == UTC
        assert (run / "prepared/index.json").read_text().startswith('{\n  "schema":')
    return run, {r["path"]: r for r in index["files"]}, code


def _book(rows: list[list[Any]], formats: dict[str, str]) -> bytes:
    book = Workbook()
    ws = book.active
    assert ws is not None
    for row in rows:
        ws.append(row)
    for ref, fmt in formats.items():
        ws[ref].number_format = fmt
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


@pytest.mark.parametrize(
    ("value", "fmt", "shown"),
    [
        (630900000, "#,##0.0,,", "630.9"),
        (6000000, "#,##0.0,,", "6"),
        (6123456, "#,##0.0,,", "6.1"),
        (6150000, "#,##0.0,,", "6.2"),
        (1234567, "#,##0,", "1235"),
        (1234567, '#,##0.00,,"M"', "1.23"),
        (1234567, "#,##0.00", "1234567"),
        (1234567, "#,##0,00", "1234567"),
        (1234567, '0 ","', "1234567"),
        (1234567, "0\\,", "1234567"),
        (1234567, "#,##0,;(#,##0,)", "1235"),
    ],
)
def test_i4_a_scaling_format_reads_as_displayed(value: int, fmt: str, shown: str) -> None:
    assert cell_text(value, fmt) == shown


@given(n=st.integers(0, 10**6), k=st.integers(0, 3))
def test_i4_xlsx_and_csv_of_the_same_displayed_sheet_agree(n: int, k: int) -> None:
    displayed = Decimal(n).scaleb(-k)
    fmt = "#,##0" + ("." + "0" * k if k else "") + ",,"
    assert Decimal(cell_text(n * 10 ** (6 - k), fmt)) == displayed


def test_i4_the_index_names_the_cells_read_as_displayed(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    rows = [["Total NAV", 630900000], ["Other", 5]]
    data = _book(rows, {"B1": "#,##0.0,,"})
    run, index, _ = _prep(tmp_path, {"w.xlsx": data})
    assert index["w.xlsx"]["display_scaled"] == {"Sheet": ["B1"]}
    assert "Total NAV,630.9\nOther,5\n" in (run / index["w.xlsx"]["prepared"]).read_text()
    assert "display-scaled (read as shown): w.xlsx sheet Sheet: B1" in capsys.readouterr().out


def test_i7_percent_cells_integer_and_quoted(tmp_path: Path) -> None:
    assert cell_text(1, "0%") == "100%"
    assert cell_text(0.5, '0.0 "%"') == "0.5"
    assert cell_text(0.5, "0.0\\%") == "0.5"
    assert cell_text(True, "0%") == "TRUE"
    assert cell_text(float("inf"), "0%") == "inf"


def _damaged_xlsx() -> dict[str, bytes]:
    good = _book([["hello", 1.5]], {})
    src = zipfile.ZipFile(io.BytesIO(good))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for i in src.infolist():
            data = src.read(i.filename)
            if i.filename.endswith("sheet1.xml"):
                data = data.replace(b"</sheetData>", b"<row><c")
            z.writestr(i, data)
    raw = bytearray(good)
    info = next(i for i in src.infolist() if i.filename.endswith("sheet1.xml"))
    start = info.header_offset + 30 + len(info.filename) + len(info.extra)
    for k in range(start + 2, start + 12):
        raw[k] ^= 0xFF
    return {"malformed.xlsx": out.getvalue(), "deflate.xlsx": bytes(raw)}


def test_i5_damaged_workbooks_are_listed_not_fatal(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    _, index, code = _prep(tmp_path, {**_damaged_xlsx(), "ok.txt": b"a plain note with words\n"})
    assert code == 0 and index["ok.txt"]["prepared"]
    out = capsys.readouterr().out
    for name in ("malformed.xlsx", "deflate.xlsx"):
        assert index[name]["unreadable"].startswith("openpyxl could not read the workbook")
        assert f"unreadable: {name}: openpyxl could not read" in out


def _surrogate_pdf() -> bytes:
    cmap = (
        b"/CIDInit /ProcSet findresource begin\n12 dict begin\nbegincmap\n/CMapName /X def\n"
        b"/CMapType 2 def\n1 begincodespacerange\n<00> <FF>\nendcodespacerange\n"
        b"1 beginbfchar\n<41> <D800>\nendbfchar\nendcmap\n"
        b"CMapName currentdict /CMap defineresource pop\nend\nend\n"
    )
    content = b"BT /F1 12 Tf 72 720 Td (Revenue grew strongly in the year AAAA) Tj ET"
    page = b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        page + b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /ToUnicode 6 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(cmap) + cmap + b"\nendstream",
    ]
    out, offs = b"%PDF-1.4\n", []
    for n, b in enumerate(objs, 1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % n + b + b"\nendobj\n"
    xref = b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    xref += b"".join(b"%010d 00000 n \n" % o for o in offs)
    trailer = b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
    return out + xref.replace(b"\n\n", b"\n") + trailer % (len(objs) + 1, len(out))


def test_i5_lone_surrogates_from_pdf_text_are_dropped_and_counted(tmp_path: Path) -> None:
    run, index, code = _prep(tmp_path, {"a.pdf": _surrogate_pdf(), "ok.txt": b"fine text here\n"})
    record = index["a.pdf"]
    assert code == 0 and record["unreadable"] is None and record["dropped_surrogates"] == 4
    text = (run / record["prepared"]).read_text(encoding="utf-8")
    assert (
        "Revenue grew strongly in the year" in text and index["ok.txt"]["dropped_surrogates"] == 0
    )


def test_i5_the_prepared_text_is_written_atomically(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def fail(*_: object) -> None:
        raise OSError("disk full")

    monkeypatch.setattr("rbotlib.inventory.os.replace", fail)
    run, _, code = _prep(tmp_path, {"a.txt": b"some text\n"})
    assert code == 1 and not list((run / "prepared").iterdir())


def test_i5_write_atomic_uses_a_unique_temp_name_and_syncs(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    seen: list[str] = []
    real = Path.unlink

    def spy(self: Path, missing_ok: bool = False) -> None:
        seen.append(self.name)
        real(self, missing_ok=missing_ok)

    synced: list[int] = []
    monkeypatch.setattr("rbotlib.inventory.os.fsync", synced.append)
    target = tmp_path / "x.json"
    write_atomic(target, "one")
    write_atomic(target, "two")
    assert target.read_text() == "two" and len(synced) == 2 and list(tmp_path.iterdir()) == [target]
    monkeypatch.setattr(Path, "unlink", spy)
    monkeypatch.setattr("rbotlib.inventory.os.replace", lambda *_: (_ for _ in ()).throw(OSError))
    with pytest.raises(OSError, match=""):
        write_atomic(target, "three")
    assert len(seen) == 1 and re.fullmatch(r"x\.json\..+\.tmp", seen[0])


BREAKS = ["\r", "\x0b", "\x0c", "\x1c", "\x1d", "\x1e", "\x85", "\u2028", "\u2029", "\r\n", "\n"]


@pytest.mark.parametrize("sep", BREAKS)
def test_i6_no_line_break_kind_can_forge_a_page_anchor(sep: str) -> None:
    text = anchored_text([f"Real one.{sep}<<< page 2 >>>\nForged", "Real two"])
    assert list(pages(text)) == [1, 2] and "Forged" in pages(text)[1]
    assert "Forged" not in pages(text)[2]


@pytest.mark.parametrize("sep", BREAKS[:-2])
def test_i6_validate_splits_pages_on_newline_only(sep: str) -> None:
    assert list(pages(f"<<< page 1 >>>\nbody{sep}<<< page 2 >>>\nmore\n")) == [1]


def test_i6_a_txt_file_with_a_carriage_return_anchor_stays_one_page(tmp_path: Path) -> None:
    run, index, _ = _prep(tmp_path, {"s.txt": b"One.\r<<< page 2 >>>\nForged\x0cReal two"})
    assert list(pages((run / index["s.txt"]["prepared"]).read_text())) == [1, 2]


def test_m1_a_non_string_prepared_is_a_named_error(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    run, _, _ = _prep(tmp_path, {"a.txt": b"alpha text body\n", "b.txt": b"bravo text body\n"})
    path = run / "prepared/index.json"
    data = json.loads(path.read_text())
    data["files"][0]["prepared"] = 123
    path.write_text(json.dumps(data))
    assert main(["prepare", "--run", str(run), "--only", "b.txt"]) == 1
    assert "'prepared' of 'a.txt' must be a path or null" in capsys.readouterr().out


def test_m1_only_never_deletes_another_files_text(tmp_path: Path) -> None:
    run, index, _ = _prep(tmp_path, {"a.txt": b"alpha text body\n", "b.txt": b"bravo text body\n"})
    path = run / "prepared/index.json"
    data = json.loads(path.read_text())
    data["files"][0]["prepared"] = index["b.txt"]["prepared"]
    path.write_text(json.dumps(data))
    assert main(["prepare", "--run", str(run), "--only", "a.txt"]) == 0
    assert (run / index["b.txt"]["prepared"]).exists()
    assert all((run / r["prepared"]).exists() for r in load_index(run)["files"])


def test_m1_only_with_no_index_says_to_run_prepare_first(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    run, _, _ = _prep(tmp_path, {"a.txt": b"alpha text body\n"})
    (run / "prepared/index.json").unlink()
    assert main(["prepare", "--run", str(run), "--only", "a.txt"]) == 1
    assert "run rbot prepare first" in capsys.readouterr().out
    assert not (run / "prepared/index.json").exists()


def test_m1_a_full_re_prepare_removes_orphaned_texts(tmp_path: Path) -> None:
    run, index, _ = _prep(tmp_path, {"a.txt": b"version one of the text\n", "b.txt": b"bravo\n"})
    old = index["a.txt"]["prepared"]
    (tmp_path / "pack/a.txt").write_text("version two of the text, changed\n")
    (tmp_path / "pack/b.txt").unlink()
    assert main(["inventory", "--run", str(run)]) == 0 and main(["prepare", "--run", str(run)]) == 0
    now = {r["path"]: r for r in load_index(run)["files"]}
    assert not (run / old).exists() and (run / now["a.txt"]["prepared"]).exists()
    assert sorted(p.name for p in (run / "prepared").glob("*.txt")) == [
        Path(now["a.txt"]["prepared"]).name
    ]


def test_m2_a_re_prepare_removes_portfolio_json_whose_id_is_gone(tmp_path: Path) -> None:
    pos = "Portfolio ID,{}\nStrategy,CLO\n\nIssuer ID,Issuer,Par,% NAV\nA,B,1,2%\n"
    run, _, _ = _prep(tmp_path, {"p.csv": pos.format("OLD").encode()})
    assert (run / "prepared/portfolios/OLD.json").exists()
    (tmp_path / "pack/p.csv").write_text(pos.format("NEW"))
    assert main(["inventory", "--run", str(run)]) == 0 and main(["prepare", "--run", str(run)]) == 0
    assert sorted(p.name for p in (run / "prepared/portfolios").iterdir()) == ["NEW.json"]


def _no_text_pdf(edit: Callable[[bytes], bytes] | None = None) -> bytes:
    raw = make_pdf([""])
    return edit(raw) if edit else raw


def test_m4_an_xref_repair_warning_keeps_a_scan_a_scan(tmp_path: Path) -> None:
    raw = _no_text_pdf()
    m = re.search(rb"startxref\n(\d+)", raw)
    assert m
    bad = raw[: m.start(1)] + str(int(m.group(1)) + 3).encode() + raw[m.end(1) :]
    _, index, _ = _prep(tmp_path, {"scan.pdf": bad})
    assert index["scan.pdf"]["unreadable"] is None and index["scan.pdf"]["image_read"] is True


def test_i7_a_no_text_pdf_whose_page_warns_is_unreadable_with_the_warning(tmp_path: Path) -> None:
    bad = _no_text_pdf(lambda raw: raw.replace(FONT.encode(), b"7".ljust(len(FONT))))
    _, index, _ = _prep(tmp_path, {"warn.pdf": bad})
    record = index["warn.pdf"]
    assert record["unreadable"] == (
        "no text extracted; pypdf warned: Page resources are not a dictionary: 7"
    )
    assert record["image_read"] is False


@pytest.mark.parametrize(
    ("value", "fmt", "shown"),
    [
        (6250000, "#,##0.0,,", "6.3"),
        (1234567, "0;0,", "1234567"),
        (1234567, "0.0,;0", "1234.6"),
        (-1234567, "#,##0,", "-1235"),
        (1.2345678901234567e300, "0.0,,", "1234567890123456700" + "0" * 276),
    ],
)
def test_i7_scaling_edge_cases(value: float, fmt: str, shown: str) -> None:
    assert cell_text(value, fmt) == shown


def test_i7_a_false_cell_reads_false() -> None:
    assert clean_cell(False) == "FALSE" and clean_cell(True) == "TRUE"


def test_i7_anchored_text_is_exact() -> None:
    assert anchored_text(["a\rb\n<<< page 2 >>>", "", "c  \n"]) == (
        "<<< page 1 >>>\na\nb\n\\<<< page 2 >>>\n<<< page 2 >>>\n<<< page 3 >>>\nc\n"
    )


def test_i7_dropped_surrogates_leave_nothing_behind(tmp_path: Path) -> None:
    run, index, _ = _prep(tmp_path, {"a.pdf": _surrogate_pdf()})
    text = (run / index["a.pdf"]["prepared"]).read_text(encoding="utf-8")
    assert text == "<<< page 1 >>>\nRevenue grew strongly in the year\n"


def test_i7_a_hidden_first_sheet_is_noted_and_the_visible_one_still_read(tmp_path: Path) -> None:
    book = Workbook()
    hidden = book.active
    assert hidden is not None
    hidden.title = "Old"
    hidden.append(["Portfolio ID", "OLD"])
    hidden.sheet_state = "hidden"
    ws = book.create_sheet("Pos")
    for row in [
        ["Portfolio ID", "P1"],
        ["Strategy", "CLO"],
        [],
        ["Issuer ID", "Issuer", "Par", "% NAV"],
        ["A", "B", 1, "2%"],
    ]:
        ws.append(row)
    out = io.BytesIO()
    book.save(out)
    run, _, _ = _prep(tmp_path, {"w.xlsx": out.getvalue()})
    index = load_index(run)
    assert [p["portfolio_id"] for p in index["portfolios"]] == ["P1"]
    assert index["portfolio_notes"] == ["w.xlsx sheet Old: hidden sheet not read for portfolios"]


@pytest.mark.parametrize(("chars", "scan"), [(19, True), (20, False)])
def test_i7_a_pdf_with_exactly_twenty_characters_has_text(
    tmp_path: Path, chars: int, scan: bool
) -> None:
    _, index, _ = _prep(tmp_path, {"t.pdf": make_pdf(["x" * chars])})
    assert index["t.pdf"]["image_read"] is scan


def test_i7_every_distinct_damage_warning_is_listed(tmp_path: Path) -> None:
    font = FONT.encode()
    raw = make_pdf(["", ""]).replace(font, b"7".ljust(len(font)), 1)
    raw = raw.replace(font, b"8".ljust(len(font)), 1)
    _, index, _ = _prep(tmp_path, {"w.pdf": raw})
    assert index["w.pdf"]["unreadable"] == (
        "no text extracted; pypdf warned: Page resources are not a dictionary: 7; "
        "Page resources are not a dictionary: 8"
    )
