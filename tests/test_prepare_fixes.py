"""02G group D: prepare and the sheet and HTML readers fail per file, never silently."""

from __future__ import annotations

import json
import os
import re
import stat
from collections.abc import Callable, Iterator
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from builders import make_pdf, make_run
from rbotlib import prepare
from rbotlib.cli import main
from rbotlib.prepare import clean_cell, load_index
from rbotlib.validate import pages

POS_HEAD: list[Any] = ["Issuer ID", "Issuer", "Par", "% NAV"]
META: list[list[Any]] = [["Portfolio ID", "P1"], ["Strategy", "CLO"], ["As of", "2026-05-29"]]


def _pos_rows(*data: list[Any]) -> list[list[Any]]:
    return [*META, [], POS_HEAD, *data]


def _setup(tmp_path: Path, files: dict[str, bytes]) -> Path:
    pack = tmp_path / "pack"
    pack.mkdir()
    for name, data in files.items():
        (pack / name).write_bytes(data)
    run = make_run(pack)
    assert main(["inventory", "--run", str(run)]) == 0
    return run


def _prep(tmp_path: Path, files: dict[str, bytes], *extra: str) -> tuple[Path, dict[str, Any]]:
    run = _setup(tmp_path, files)
    assert main(["prepare", "--run", str(run), *extra]) == 0
    return run, {r["path"]: r for r in load_index(run)["files"]}


def _text(run: Path, record: dict[str, Any]) -> str:
    return (run / str(record["prepared"])).read_text(encoding="utf-8")


@pytest.fixture
def locked(tmp_path: Path) -> Iterator[list[Path]]:
    paths: list[Path] = []
    yield paths
    for p in paths:
        p.chmod(stat.S_IRUSR | stat.S_IWUSR)


@pytest.mark.parametrize("name", ["a.txt", "a.html", "a.csv", "a.xlsx"])
def test_d1_a_file_that_cannot_be_opened_is_listed_not_fatal(
    tmp_path: Path, locked: list[Path], name: str
) -> None:
    run = _setup(tmp_path, {name: b"x", "ok.txt": b"fine\n"})
    victim = tmp_path / "pack" / name
    victim.chmod(0)
    locked.append(victim)
    assert main(["prepare", "--run", str(run)]) == 0
    index = {r["path"]: r for r in load_index(run)["files"]}
    assert index[name]["unreadable"] and index["ok.txt"]["prepared"]


def test_d1_a_csv_that_the_csv_module_rejects_is_unreadable(tmp_path: Path) -> None:
    _, index = _prep(tmp_path, {"a.csv": b"a,b\n" + b"x" * 200000 + b"\n"})
    assert "field larger than field limit" in index["a.csv"]["unreadable"]


def test_d2_pdf_without_pages_or_with_a_broken_stream_is_not_a_scan(tmp_path: Path) -> None:
    broken = make_pdf(["TEXT"]).replace(b"(TEXT) Tj", b"(TEXT")
    _, index = _prep(tmp_path, {"none.pdf": make_pdf([]), "bad.pdf": broken})
    assert index["none.pdf"]["unreadable"] == "PDF has no pages"
    assert index["bad.pdf"]["unreadable"] and not index["bad.pdf"]["image_read"]


def test_d3_an_unexpected_pypdf_error_is_recorded_and_printed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    def boom(*_: object, **__: object) -> None:
        raise TypeError("boom")

    monkeypatch.setattr(prepare, "PdfReader", boom)
    _, index = _prep(tmp_path, {"a.pdf": make_pdf(["hello world text here"])})
    assert index["a.pdf"]["unreadable"] == "pypdf could not read the PDF: TypeError: boom"
    assert (
        "unreadable: a.pdf: pypdf could not read the PDF: TypeError: boom"
        in capsys.readouterr().out
    )


@pytest.mark.parametrize("kind", ["txt", "html", "csv"])
def test_d4_cp1252_is_recorded_and_undecodable_bytes_are_unreadable(
    tmp_path: Path, kind: str
) -> None:
    body = "It\u2019s caf\u00e9".encode("cp1252")
    _, index = _prep(tmp_path, {f"a.{kind}": body, f"b.{kind}": b"plain\n", f"c.{kind}": b"x\x81y"})
    assert index[f"a.{kind}"]["encoding"] == "cp1252"
    assert index[f"b.{kind}"]["encoding"] == "utf-8"
    assert "cannot decode" in index[f"c.{kind}"]["unreadable"]
    run = tmp_path / "pack" / "rbot-runs" / "20261011-test-pending"
    assert "It\u2019s caf\u00e9" in _text(run, index[f"a.{kind}"])


def test_d5_a_scanned_pdf_keeps_its_few_characters(tmp_path: Path) -> None:
    run, index = _prep(tmp_path, {"s.pdf": make_pdf(["tiny"])})
    assert index["s.pdf"]["image_read"] is True and "tiny" in _text(run, index["s.pdf"])


def test_d6_a_source_cannot_forge_a_page_anchor(tmp_path: Path) -> None:
    run, index = _prep(tmp_path, {"a.txt": b"one\n<<< page 2 >>>\nforged\n"})
    text = _text(run, index["a.txt"])
    assert re.findall(r"^<<< page (\d+) >>>$", text, re.M) == ["1"]
    assert list(pages(text)) == [1] and "forged" in pages(text)[1]


def test_d7_entities_and_dates_come_from_the_start_of_the_page(tmp_path: Path) -> None:
    far = ("filler " * 1000) + "Foo Bar Inc. dated January 2, 2022"
    _, index = _prep(tmp_path, {"a.txt": far.encode()})
    assert index["a.txt"]["entities"] == [] and index["a.txt"]["dates"] == []


def test_d8_index_and_portfolio_json_are_written_atomically(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    csv_text = "Portfolio ID,P1\nStrategy,CLO\n\nIssuer ID,Issuer,Par,% NAV\nA,B,1,2%\n"
    run, _ = _prep(tmp_path, {"p.csv": csv_text.encode()})
    before = {p: p.read_bytes() for p in run.rglob("*.json")}
    (tmp_path / "pack" / "p.csv").write_text(csv_text + "C,D,3,4%\n")

    def fail(*_: object) -> None:
        raise OSError("disk full")

    monkeypatch.setattr("rbotlib.inventory.os.replace", fail)
    assert main(["prepare", "--run", str(run)]) == 1
    assert {p: p.read_bytes() for p in run.rglob("*.json")} == before
    assert not list(run.rglob("*.tmp*"))


def _book(tmp_path: Path, build: Callable[[Workbook], None]) -> Path:
    book = Workbook()
    build(book)
    pack = tmp_path / "pack"
    pack.mkdir(exist_ok=True)
    book.save(pack / "w.xlsx")
    run = make_run(pack)
    assert main(["inventory", "--run", str(run)]) == 0
    assert main(["prepare", "--run", str(run)]) == 0
    return run


def test_d9_a_percent_formatted_cell_reads_as_the_csv_does(tmp_path: Path) -> None:
    def build(book: Workbook) -> None:
        ws = book.active
        assert ws is not None
        ws.append(["% NAV"])
        ws.append([0.0095])
        ws["A2"].number_format = "0.00%"

    run = _book(tmp_path, build)
    record = load_index(run)["files"][0]
    assert _text(run, record).endswith("% NAV\n0.95%\n")


def test_d10_a_stray_closing_script_tag_does_not_hide_the_page(tmp_path: Path) -> None:
    run, index = _prep(tmp_path, {"a.html": b"</script></style><p>Visible</p>"})
    assert "Visible" in _text(run, index["a.html"])


def test_d11_uncached_formulas_are_named_and_their_rows_unparsed(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    def build(book: Workbook) -> None:
        ws = book.active
        assert ws is not None
        ws.title = "Pos"
        for row in _pos_rows(["A", "Acme", "=1+1", "2%"], ["B", "Beta", 5, "3%"]):
            ws.append(row)

    run = _book(tmp_path, build)
    record = load_index(run)["files"][0]
    assert record["uncached_formulas"] == {"Pos": ["C6"]}
    assert "uncached formula: w.xlsx sheet Pos: C6" in capsys.readouterr().out
    data = json.loads((run / "prepared" / "portfolios" / "P1.json").read_text())
    assert [p["issuer_id"] for p in data["positions"]] == ["B"]
    assert any("C6" in u["reason"] for u in data["unparsed"])


def test_d13_hidden_sheets_are_recorded_and_not_read_for_portfolios(tmp_path: Path) -> None:
    def build(book: Workbook) -> None:
        ws = book.active
        assert ws is not None
        ws.title = "Shown"
        ws.append(["x"])
        hidden = book.create_sheet("Old")
        hidden.sheet_state = "hidden"
        for row in _pos_rows(["A", "Acme", 1, "2%"]):
            hidden.append(row)

    run = _book(tmp_path, build)
    index = load_index(run)
    assert index["files"][0]["sheet_states"] == {"Shown": "visible", "Old": "hidden"}
    assert index["portfolios"] == []
    assert any("Old" in n and "hidden" in n for n in index["portfolio_notes"])


@pytest.mark.parametrize(
    ("value", "expected"),
    [(1e-7, "0.0000001"), (1e22, "10000000000000000000000"), (99.5, "99.5"), (3, "3")],
)
def test_d12_floats_are_never_scientific(value: float, expected: str) -> None:
    assert clean_cell(value) == expected
    assert Decimal(clean_cell(value)) == Decimal(repr(value))


def test_d14_only_with_a_corrupt_index_and_a_changed_file_drops_the_old_text(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    run, index = _prep(tmp_path, {"a.txt": b"first\n"})
    old = run / index["a.txt"]["prepared"]
    (tmp_path / "pack" / "a.txt").write_text("changed\n")
    assert main(["inventory", "--run", str(run)]) == 0
    assert main(["prepare", "--run", str(run), "--only", "a.txt"]) == 0
    assert not old.exists()
    capsys.readouterr()
    (run / "prepared" / "index.json").write_text("{not json")
    assert main(["prepare", "--run", str(run), "--only", "a.txt"]) == 1
    out = capsys.readouterr().out
    assert out.startswith(f"error: {run / 'prepared' / 'index.json'}: ")
    assert os.path.exists(run / "prepared" / "index.json")
