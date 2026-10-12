from __future__ import annotations

import json
from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from conftest import CORPUS
from rbotlib.corpus import repo_root
from rbotlib.prepare import load_index
from rbotlib.validate import excerpt_found

SYNTHETIC = repo_root(CORPUS) / "tests" / "corpus" / "synthetic"
PRIMARY = "czr-2026q2-primary"


def _first_line(path: Path) -> str:
    if path.suffix == ".xlsx":
        book = load_workbook(path, read_only=True)
        first = next(book.worksheets[0].iter_rows(max_row=1, values_only=True))
        book.close()
        return str(first[0])
    return path.read_text(encoding="utf-8").splitlines()[0]


def _committed(pack_path: str, f: dict[str, Any]) -> str:
    """Where a SYNTHETIC file is committed: its own pack, or the pack its `external` names
    (amendment A12: the portfolio packs share the primary pack's Test CLO I files)."""
    root, _, rest = str(f.get("external", "")).partition("/")
    return rest if root == "synthetic" else f"{pack_path}/{f['path']}"


def test_every_authored_file_is_labelled_and_registered() -> None:
    committed = sorted(p for p in SYNTHETIC.rglob("*") if p.is_file())
    registered = {
        _committed(pack.path, f)
        for pack in CORPUS.packs
        for f in pack.data["files"]
        if f.get("synthetic")
    }
    assert {p.relative_to(SYNTHETIC).as_posix() for p in committed} == registered
    for path in committed:
        assert "SYNTHETIC" in path.name
        if path.suffix != ".png":  # the screenshot carries its label in the picture
            assert _first_line(path).startswith("SYNTHETIC"), path.name


def _records(run: Path) -> dict[str, dict[str, Any]]:
    return {r["path"]: r for r in load_index(run)["files"]}


@pytest.mark.corpus
def test_primary_pack_traps(prepared_run: Callable[[str], Path]) -> None:
    run = prepared_run(PRIMARY)
    inventory = json.loads((run / "inventory.json").read_text(encoding="utf-8"))
    digests = {f["path"]: f["sha256"] for f in inventory["files"]}
    assert (
        digests["TRACE_12769GAC4_last_trade.txt"]
        == digests["CZR_FINRA_TRACE_12769GAC4_2026-10-02.txt"]
    )
    records = _records(run)
    assert records["SYNTHETIC_Levels_Screenshot.png"]["image_read"] is True
    vmo2 = records["VMED-O2-Quarterly-Bond-Report-Q1-2026.pdf"]
    assert vmo2["pages"] == 47 and vmo2["entities"] == ["VMED O2 UK HOLDINGS LIMITED"]
    assert records["SYNTHETIC_Sector_RV_Gaming.xlsx"]["sheets"] == ["Gaming"]
    levels = (run / records["SYNTHETIC_Bloomberg_Levels_2026-10-02.csv"]["prepared"]).read_text()
    assert "CZR TL B 2030,BL3912045,,1L Term Loan,99.375,99.750,99.563,,252,Ba3,BB\n" in levels


@pytest.mark.corpus
def test_primary_pack_portfolio_json(prepared_run: Callable[[str], Path]) -> None:
    run = prepared_run(PRIMARY)
    assert load_index(run)["portfolios"] == [
        {"portfolio_id": "TEST_CLO_I", "json": "prepared/portfolios/TEST_CLO_I.json"}
    ]
    data = json.loads((run / "prepared" / "portfolios" / "TEST_CLO_I.json").read_text())
    assert (data["strategy"], data["nav"], data["par"], data["unit"]) == (
        "CLO",
        "630.9",
        "634.0",
        "MILLIONS",
    )
    assert len(data["positions"]) == 10
    assert not [p for p in data["positions"] if p["issuer_id"] == "CZR"]
    assert [lim["limit_id"] for lim in data["limits"]] == [
        "C-01", "C-02", "C-04", "C-05", "C-06", "C-07", "C-08", "C-10", "C-11", "C-17",
    ]  # fmt: skip
    unparsed = [u["row"].rsplit(" ", 1)[-1] for u in data["unparsed"]]
    assert unparsed == ["9", "15", "18", "19", "20", "21", "22"]  # C-03, C-09, C-12 to C-16


@pytest.mark.corpus
@pytest.mark.parametrize("pack_id", [p.id for p in CORPUS.packs if p.data["expectations"]])
def test_expectations_anchor_in_their_sources(
    pack_id: str, prepared_run: Callable[[str], Path]
) -> None:
    expected = json.loads(
        (CORPUS.location.parent / CORPUS.pack(pack_id).data["expectations"]).read_text()
    )
    assert expected["pack"] == pack_id
    run = prepared_run(pack_id)
    records = _records(run)
    for item in expected["values"]:
        assert {"module", "table", "where", "value"} <= set(item), item
        if "source" in item:
            assert item.get("locator"), item
            text = (run / records[item["source"]]["prepared"]).read_text(encoding="utf-8")
            assert excerpt_found(text, item["excerpt"]), item
        else:
            assert "rule" in item, item


def test_the_synthetic_notes_state_their_dates_and_the_expected_downgrade() -> None:
    folder = SYNTHETIC / PRIMARY
    for name in ("Positions", "Limits"):
        text = (folder / f"SYNTHETIC_TEST_CLO_I_{name}_2026-10-02.csv").read_text(encoding="utf-8")
        note = next(line for line in text.splitlines() if line.startswith("Note,"))
        assert "as of 2026-05-29" in note
    terms = (folder / "SYNTHETIC_CZR_2026_TLB_Term_Sheet.md").read_text(encoding="utf-8")
    assert "An S&P downgrade of the existing first-lien to BB- is expected at closing." in terms
