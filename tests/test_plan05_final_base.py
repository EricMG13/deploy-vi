"""Real final-base producer/consumer controls; fixture evidence is not model qualification."""

import json
import shutil
from collections.abc import Callable
from decimal import Decimal
from pathlib import Path

import ab_evidence
import pytest
from openpyxl import load_workbook
from test_cp1 import _model
from test_fit import _run
from test_p05_review_followup import _a_evidence

import ab_grade
from conftest import FIXTURES
from rbotlib import validate
from rbotlib.cli import main
from rbotlib.handoff import parse


def test_cp5_closed_ltm_finding_matches_the_actual_model(tmp_path: Path) -> None:
    data = _model(tmp_path, "CP-1", "CP-1D")
    assert data["house"]["structuring_ebitda"]["value"] == "3742"
    assert data["house"]["ssim_ebitda"]["value"] == "3592"
    cp5 = parse((FIXTURES / "handoffs/CP-5.md").read_text("utf-8"))
    table = cp5.tables["cp5.findings"]
    finding = next(
        dict(zip(table.columns, row, strict=True)) for row in table.rows if row[0] == "F-1"
    )
    assert finding["status"] == "Fixed"
    assert "3,742" in finding["description"]
    assert not any(
        "leaves out" in cell for row in cp5.tables["cp5.limitations"].rows for cell in row
    )
    statuses = cp5.tables["cp5.module_status"]
    model = next(
        dict(zip(statuses.columns, row, strict=True)) for row in statuses.rows if row[0] == "MODEL"
    )
    assert model["traceability"] == "Traceable"
    assert "misses" not in model["note"]


@pytest.mark.parametrize(
    ("status", "verdict", "ssim"),
    [
        ("UNREALIZED", "Challenged", "3592"),
        ("REALIZED", "Rejected", "3592"),
        ("REALIZED", "Supported", "3742"),
    ],
)
def test_plan05_current_ltm_judgment_reaches_all_model_outputs(
    tmp_path: Path, status: str, verdict: str, ssim: str
) -> None:
    """REALIZED/Supported is a hypothetical guard control, not an achieved-source claim."""
    run = _run(tmp_path, None)
    text = (FIXTURES / "handoffs/CP-1D.md").read_text("utf-8")
    table = parse(text).tables["cp1d.addback_assessment"]
    old = next(row for row in table.rows if row[0] == "run_rate_synergies")
    new = list(old)
    new[table.columns.index("house_realization_status")] = status
    new[table.columns.index("verdict")] = verdict
    new[table.columns.index("rationale")] = "HYPOTHETICAL current-LTM judgment control"
    text = text.replace("| " + " | ".join(old) + " |", "| " + " | ".join(new) + " |")
    (run / "handoffs/CP-1D.md").write_text(text, "utf-8")
    assert main(["model", "--run", str(run)]) == 0
    data = json.loads((run / "model.json").read_text("utf-8"))
    assert data["house"]["structuring_ebitda"]["value"] == "3742"
    assert data["house"]["ssim_ebitda"]["value"] == ssim
    model = parse((run / "handoffs/MODEL.md").read_text("utf-8"))
    rows = model.tables["model.house"].rows
    for item, expected in (("Structuring EBITDA", "3742"), ("SSIM EBITDA", ssim)):
        assert next(row for row in rows if row[0] == item)[1] == expected
    book = load_workbook(run / "model.xlsx")
    for item, expected in (("Structuring EBITDA", "3742"), ("SSIM EBITDA", ssim)):
        row = next(row for row in book["House"].iter_rows(values_only=True) if row[0] == item)
        assert Decimal(str(row[1])) == Decimal(expected)
    book.close()


@pytest.mark.parametrize("side", ["A", "B"])
def test_both_sides_refuse_a_genuine_global_prepared_index_failure(
    tmp_path: Path, side: str
) -> None:
    text, run, mapping = _a_evidence(tmp_path)
    assert vars(ab_evidence)["run_findings"] is validate.run_findings
    assert vars(ab_grade)["run_findings"] is validate.run_findings
    (run / "prepared/index.json").write_text("{broken", "utf-8")
    with pytest.raises(ab_grade.ABError, match=r"run_error.*prepared/index.json"):
        if side == "A":
            ab_grade.measure_a(text, run, mapping)
        else:
            ab_grade.measure("CP-5", (FIXTURES / "handoffs/CP-5.md").read_text("utf-8"), run, None)


@pytest.mark.parametrize(
    ("native", "source"),
    [
        ("5 million for the quarter", "-$5 million for the quarter"),
        ("$5 million for the quarter", "- $5 million for the quarter"),
        ("$5 million for the quarter", "($5 million for the quarter)"),
    ],
)
def test_native_measurement_uses_the_actual_sign_currency_anchor(
    tmp_path: Path, native: str, source: str
) -> None:
    text, run, path = _a_evidence(tmp_path)
    data = json.loads(path.read_text("utf-8"))
    data["citations"][0]["excerpt"] = native
    path.write_text(json.dumps(data), "utf-8")
    text = text.replace("Revenue 100.", native)
    (run / "prepared/SYNTHETIC_annual.txt").write_text(
        "<<< page 1 >>>\nSYNTHETIC " + source + "\n", "utf-8"
    )
    assert ab_grade.measure_a(text, run, path)["C"] == "0/1"


@pytest.mark.corpus
def test_real_prepared_model_cp5_measure_fit_and_cp6_chain(
    prepared_run: Callable[[str], Path], tmp_path: Path
) -> None:
    run = tmp_path / "pack/rbot-runs/20261011-czr-full-credit-assessment"
    shutil.copytree(prepared_run("czr-2026q2-primary"), run)
    (run / "handoffs").mkdir(exist_ok=True)
    for module in ("CP-0", "CP-1", "CP-1A", "CP-1B", "CP-1D", "CP-2", "CP-3", "CP-5", "CP-6"):
        shutil.copy(FIXTURES / f"handoffs/{module}.md", run / f"handoffs/{module}.md")
    assert main(["brief", "--run", str(run), "--from", "handoffs/CP-0.md"]) == 0
    assert main(["model", "--run", str(run)]) == 0
    data = json.loads((run / "model.json").read_text("utf-8"))
    assert (
        data["house"]["structuring_ebitda"]["value"],
        data["house"]["ssim_ebitda"]["value"],
    ) == ("3742", "3592")
    cp5 = (run / "handoffs/CP-5.md").read_text("utf-8")
    measured = ab_grade.measure("CP-5", cp5, run, None)
    assert measured["C"] == "2/2" and measured["validate_exit"] == 0
    poor = cp5.replace(
        "| Earnings and interest | Run-rate cost synergies", "| p. 999 | Run-rate cost synergies"
    )
    assert poor != cp5
    measured = ab_grade.measure("CP-5", poor, run, None)
    assert measured["C"] == "1/2" and measured["validate_exit"] == 2
    assert main(["fit", "--run", str(run)]) == 0
    fits = {
        tuple(row[:3]): row
        for row in parse((run / "handoffs/FIT.md").read_text("utf-8"))
        .tables["snapshot.portfolio_fit"]
        .rows
    }
    cp6 = parse((run / "handoffs/CP-6.md").read_text("utf-8"))
    cap_rows = cp6.tables["cp6.portfolio_recommendation"]
    for row in cap_rows.rows:
        item = dict(zip(cap_rows.columns, row, strict=True))
        fit = fits[(item["portfolio_id"], item["strategy"], item["instrument"])]
        assert item["cap"] == fit[11] and item["fit_verdict"] == fit[10]
    measured = ab_grade.measure(
        "CP-6",
        (run / "handoffs/CP-6.md").read_text("utf-8"),
        run,
        FIXTURES.parent / "corpus/expectations/czr-2026q2-primary.json",
    )
    assert measured["F"] == "15/15" and measured["validate_exit"] == 0
