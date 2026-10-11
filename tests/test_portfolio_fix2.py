"""02G fix round 2: one table per sheet, ambiguous units refused, bounded numbers."""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from openpyxl import Workbook
from test_portfolio_fixes import AMOUNT, LIM, LIM_HEAD, POS, _item, _json

from builders import make_run
from rbotlib.cli import main
from rbotlib.portfolio import (
    Portfolio,
    Table,
    _limit_row,
    _position_values,
    number,
    parse_limit,
    write_all,
)
from rbotlib.rules import load_rules

ONE = "more than one table in this sheet; put each portfolio table on its own sheet or file"
PH = "Issuer ID,Issuer,Par,% NAV\n"
SYNTH = Path(__file__).parent / "corpus/synthetic/czr-2026q2-primary"


def _meta(pid: str) -> str:
    return f"Portfolio ID,{pid}\nStrategy,CLO\nAs of,2026-05-29\nTotal NAV,100\nTotal Par,100\n"


def _lim(value: str) -> str:
    return f"C-01,Single Name,Max x,≤ {value}% NAV,Hard,0.4%\n"


_A = PH + "A,Acme,1,1%\n"
_B = PH + "B,Beta,2,2%\n"
STACKED = {
    "stacked, own meta": _meta("P1") + "\n" + _A + "\n" + _meta("P2") + "\n" + _B,
    "no blank line": _meta("P1") + _A + _meta("P2") + _B + LIM_HEAD + _lim("2.5"),
    "blank lines": _meta("P1") + _A + "\n\n" + _meta("P2") + _B + "\n\n" + LIM_HEAD + _lim("2.5"),
    "source row": _meta("P1") + _A + "\nPortfolio ID,P2\nSource,x\n" + LIM_HEAD + _lim("9.9"),
    "note row": _meta("P1") + _A + "\nPortfolio ID,P2\nNote,y\n" + LIM_HEAD + _lim("9.9"),
    "title row": _meta("P1") + _A + "\nPortfolio ID,P2\nCompliance monitor,T,2026\n" + LIM_HEAD,
    "loans then bonds": _meta("P1") + "\nLoans\n" + _A + "\nBonds\n" + PH + "C,Czr,5,5%\n",
    "repeated header": _meta("P1") + "\n" + _A + PH + "C,Czr,5,5%\n",
    "limits without id": _meta("P1") + _A + "\n" + _meta("P2") + _B + "\nAs of,x\n" + LIM_HEAD,
    "one positions and one limits": POS.format(unit="") + "\n" + LIM_HEAD + _lim("2.5"),
    "stacked Unit": POS.format(unit="Unit,USD\n") + "\nUnit,MILLIONS\n" + LIM_HEAD + AMOUNT,
}


@pytest.mark.parametrize("text", list(STACKED.values()), ids=list(STACKED))
def test_n1_a_sheet_with_more_than_one_table_is_parsed_into_no_portfolio(
    tmp_path: Path, text: str
) -> None:
    data, notes = _json(tmp_path, {"one.csv": text})
    assert data == {}
    assert len(notes) == 1 and notes[0].startswith("one.csv: " + ONE)


def test_n1_two_portfolios_in_one_synthetic_file_are_refused(tmp_path: Path) -> None:
    lim = (SYNTH / "SYNTHETIC_TEST_CLO_I_Limits_2026-10-02.csv").read_text(encoding="utf-8")
    pos = (SYNTH / "SYNTHETIC_TEST_CLO_I_Positions_2026-10-02.csv").read_text(encoding="utf-8")
    pos2 = pos.replace("Portfolio ID,TEST_CLO_I", "Portfolio ID,TEST_CLO_II")
    data, notes = _json(tmp_path, {"combined.csv": lim + "\n" + pos2})
    assert data == {} and ONE in notes[0]


def test_n1_the_other_files_of_the_pack_are_still_read(tmp_path: Path) -> None:
    pos = POS.format(unit="")
    data, notes = _json(tmp_path, {"two.csv": STACKED["repeated header"], "ok.csv": pos})
    assert list(data) == ["P1"] and data["P1"]["positions_file"] == "ok.csv"
    assert len(notes) == 1 and "two.csv" in notes[0]


@pytest.mark.parametrize("cells", ["USD,millions", "USD,000s", "MILLIONS,USD"])
@pytest.mark.parametrize("where", ["positions", "limits"])
@pytest.mark.parametrize("order", [("a_pos.csv", "b_lim.csv"), ("b_pos.csv", "a_lim.csv")])
def test_c1_a_unit_row_with_several_values_makes_every_amount_limit_unparsed(
    tmp_path: Path, cells: str, where: str, order: tuple[str, str]
) -> None:
    row = f"Unit,{cells}"
    pos = POS.format(unit=row + "\n" if where == "positions" else "")
    lim = LIM.format(unit=row if where == "limits" else "") + AMOUNT
    d = _json(tmp_path, {order[0]: pos, order[1]: lim})[0]["P1"]
    assert d["limits"] == []
    reasons = [u["reason"] for u in d["unparsed"] if u["row"].startswith(order[1])]
    assert len(reasons) == 1 and reasons[0].startswith("ambiguous Unit row: ")
    assert all(cell in reasons[0] for cell in cells.split(","))
    assert d["unit"] == cells.replace(",", ", ")


def _book(unit: str) -> bytes:
    book = Workbook()
    ws = book.active
    assert ws is not None
    ws.title = "Pos"
    head: list[object] = ["Issuer ID", "Issuer", "Par", "% NAV"]
    rows: list[list[object]] = [["Portfolio ID", "P1"], ["Strategy", "CLO"]]
    rows += [["As of", "2026-05-29"], ["Unit", unit], ["Total NAV", 630900000]]
    rows += [["Total Par", 634000000], [], head]
    for r in rows:
        ws.append(r)
    ws.append(["CZR", "Caesars", 6000000, 0.0095])
    for ref in ("B5", "B6", "C9"):
        ws[ref].number_format = "#,##0.0,,"
    lim = book.create_sheet("Lim")
    for line in [["Portfolio ID", "P1"], [], LIM_HEAD.strip().split(",")]:
        lim.append(line)
    lim.append(["C-01", "Single Name", "Max single issuer", "≤ $25,000,000", "Hard", 6000000])
    lim["F4"].number_format = "#,##0.0,,"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


@pytest.mark.parametrize("unit", ["USD", "MILLIONS"])
def test_n2_a_display_scaled_cell_makes_the_amount_unit_uncertain(
    tmp_path: Path, unit: str
) -> None:
    pack = tmp_path / "pack"
    pack.mkdir()
    (pack / "book.xlsx").write_bytes(_book(unit))
    run = make_run(pack)
    assert main(["inventory", "--run", str(run)]) == 0 and main(["prepare", "--run", str(run)]) == 0
    d = json.loads((run / "prepared/portfolios/P1.json").read_text(encoding="utf-8"))
    assert d["limits"] == []
    if unit == "USD":
        why = "values are display-scaled; the amount unit is not certain"
        assert [u["reason"] for u in d["unparsed"] if "Lim" in u["row"]] == [why]


@pytest.mark.parametrize("total", ["630,928,502", "1,2,3", "5,6"])
def test_n4_a_total_row_with_more_than_two_cells_is_unparsed_never_its_first_group(
    tmp_path: Path, total: str
) -> None:
    d = _json(tmp_path, {"p.csv": POS.format(unit="").replace("630.9", total)})[0]["P1"]
    assert d["nav"] == "" and d["par"] == "634.0"
    rows = [u for u in d["unparsed"] if u["row"] == "Total NAV"]
    assert len(rows) == 1 and "more than one value" in rows[0]["reason"]


def test_n4_a_quoted_total_with_thousands_groups_is_one_cell(tmp_path: Path) -> None:
    d = _json(tmp_path, {"p.csv": POS.format(unit="").replace("630.9", '"630,928,502"')})[0]["P1"]
    assert d["nav"] == "630928502"


def test_n5_numbers_longer_than_thirty_significant_digits_are_not_numbers() -> None:
    assert number("1" * 31) is None and number("0." + "1" * 31) is None
    assert number("1" * 30) == ("1" * 30, "")
    assert number("0.000" + "1" * 30) == ("0.000" + "1" * 30, "")
    long = "1" * 31
    assert parse_limit(f"≤ {long}% NAV") == f"limit '≤ {long}% NAV' has more than 30 digits"
    assert parse_limit(f"≤ ${long}") == f"limit '≤ ${long}' has more than 30 digits"


def test_n5_a_count_current_value_of_29_digits_and_a_fraction_is_unparsed() -> None:
    row = _item("≤ 25", "1" * 29 + ".5", "Max number of issuers")
    result = _limit_row(Table("t", [], 0, {}), 5, row, [])
    assert result == f"count current value '{'1' * 29}.5' is not a whole number"


def test_n5_the_csv_case_runs_to_the_end(tmp_path: Path) -> None:
    lim = "Portfolio ID,P1\n\n" + LIM_HEAD
    lim += "C-01,Count,Max number of issuers,≤ 25,Hard," + "1" * 29 + ".5\n"
    d = _json(tmp_path, {"l.csv": lim})[0]["P1"]
    assert d["limits"] == [] and "not a whole number" in d["unparsed"][0]["reason"]


@pytest.mark.parametrize("text", ["1,500%", "7,500%", "1,500.5%", "$1,500%"])
def test_i3_a_percent_with_a_thousands_group_is_ambiguous(text: str) -> None:
    assert number(text) is None
    why = "percent with a thousands separator is ambiguous"
    assert parse_limit(f"≤ {text} NAV") == f"limit '≤ {text} NAV': {why}"
    row = _item("≤ 2.5% NAV", text)
    assert _limit_row(Table("t", [], 0, {}), 5, row, []) == (
        f"current value {text!r} is not a number ({why})"
    )


def test_i3_a_thousands_group_is_still_fine_without_a_percent() -> None:
    assert number("1,500") == ("1500", "") and number("1,500 bps") == ("1500", "bps")
    assert parse_limit("≤ $1,500") == ("1500", "amount")


def test_n6_the_stale_sweep_compares_names_case_insensitively(tmp_path: Path) -> None:
    folder = tmp_path / "prepared" / "portfolios"
    folder.mkdir(parents=True)
    (folder / "p1.json").write_text("{}", encoding="utf-8")
    (folder / "old.json").write_text("{}", encoding="utf-8")
    meta_only = Portfolio(meta={"portfolio_id": "P1"})
    written = write_all(tmp_path, {"P1": meta_only}, load_rules())
    assert written == [{"portfolio_id": "P1", "json": "prepared/portfolios/P1.json"}]
    assert (folder / "P1.json").is_file() and not (folder / "old.json").exists()
    assert json.loads((folder / "P1.json").read_text(encoding="utf-8"))["portfolio_id"] == "P1"


def test_n5_i3_the_reason_of_a_position_names_the_refused_number() -> None:
    long, pct = "1" * 31, "1,500%"
    assert _position_values({"par": long, "pct_nav": "1%"}) == (
        "par or % NAV is not a number (more than 30 digits)"
    )
    assert _position_values({"par": "5", "pct_nav": pct}) == (
        "par or % NAV is not a number (percent with a thousands separator is ambiguous)"
    )
    assert _position_values({"par": "5", "pct_nav": "x"}) == "par or % NAV is not a number"
