"""model.xlsx's guards (up-front review of the non-money modules,
2026-10-11): text from CP-1 never becomes a formula, and the House and
Snapshot sheets hold the Model's exact figures with number formats."""

from __future__ import annotations

import zipfile
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import cast

import pytest
from credit_model_fixture import example_registers
from model_handoffs import cp1_handoff, run_brief, run_folder
from openpyxl import load_workbook
from run_cases import _model, _run

from rbotlib.cli import main
from rbotlib.model.xlsx import FORMATS


def _sheet_xml(run: Path) -> list[str]:
    with zipfile.ZipFile(run / "model.xlsx") as book:
        names = [n for n in book.namelist() if n.startswith("xl/worksheets/sheet")]
        return [book.read(n).decode("utf-8") for n in names]


def _excerpt_changed(text: str, excerpt: str) -> str:
    lines = text.splitlines()
    at = next(n for n, line in enumerate(lines) if line.startswith("| C3 |"))
    lines[at] = lines[at].rsplit("|", 2)[0] + f"| {excerpt.replace('|', chr(92) + '|')} |"
    return "\n".join(lines)


@pytest.mark.parametrize("lead", ["=", "+", "-", "@"])
def test_text_from_cp1_never_becomes_a_formula(tmp_path: Path, lead: str) -> None:
    # I4: openpyxl stores any string starting with `=` as a formula.
    text = cp1_handoff(example_registers())
    text = text.replace("Las Vegas hotel occupancy", f'{lead}HYPERLINK("http://x.example","click")')
    text = _excerpt_changed(text, f"{lead}cmd|' /C calc'!A0")
    run = run_folder(tmp_path, text, run_brief())
    assert main(["model", "--run", str(run)]) == 0
    for xml in _sheet_xml(run):
        assert "<f>" not in xml and "<f " not in xml
    book = load_workbook(BytesIO((run / "model.xlsx").read_bytes()))
    texts = [c.value for s in book for r in s.iter_rows() for c in r if isinstance(c.value, str)]
    assert f'{lead}HYPERLINK("http://x.example","click")' in texts
    assert any(t.startswith(f"{lead}cmd") for t in texts)
    assert all(c.data_type != "f" for s in book for r in s.iter_rows() for c in r)


def test_the_house_sheet_holds_exact_values_with_number_formats(tmp_path: Path) -> None:
    # I6: the sheet held the display text ("4134", "98.43x").
    run = _run(tmp_path)
    assert main(["model", "--run", str(run)]) == 0
    house = load_workbook(BytesIO((run / "model.xlsx").read_bytes()))["House"]
    header = [c.value for c in house[1]]
    assert header == [
        "item", "value", "unit_label", "column_id", "formula", "components", "reason"
    ]  # fmt: skip
    rows = {r[0].value: r for r in house.iter_rows(min_row=2)}
    ssim = rows["SSIM EBITDA"][1]
    assert (ssim.value, ssim.number_format) == (4094, FORMATS["MONEY"])
    multiple = rows["SSIM Multiple"][1]
    exact = Decimal(str(_model(run)["house"]["ssim_multiple"]["value"]))
    assert multiple.number_format == FORMATS["MULTIPLE"]
    assert abs(Decimal(str(multiple.value)) - exact) < Decimal("1e-12")
    assert str(rows["SSIM EBITDA"][5].value).startswith("adjusted_ebitda")
    assert rows["Adjusted debt"][5].value == "total_debt"


def test_the_snapshot_sheet_holds_values_and_blank_cells_with_a_passport_reason(
    tmp_path: Path,
) -> None:
    run = _run(tmp_path)
    assert main(["model", "--run", str(run)]) == 0
    book = load_workbook(BytesIO((run / "model.xlsx").read_bytes()))
    snapshot = book["Snapshot"]
    header = [c.value for c in snapshot[1]]
    rows = {r[2].value: r for r in snapshot.iter_rows(min_row=2)}
    revenue = rows["revenue"][header.index("fy")]
    assert (revenue.value, revenue.number_format) == (11519, FORMATS["MONEY"])
    margin = rows["metric:gross_margin"][header.index("fy")]
    assert margin.number_format == FORMATS["PERCENT"] and 0 < Decimal(str(margin.value)) < 1
    growth = rows["revenue_growth"][header.index("fy_2")]
    assert growth.value is None
    passport = [[c.value for c in r] for r in book["Passport"].iter_rows(min_row=2)]
    assert ["Snapshot", growth.row, growth.column, "revenue_growth", "fy_2"] == next(
        p[:5]
        for p in passport
        if p[0] == "Snapshot" and p[3] == "revenue_growth" and p[4] == "fy_2"
    )


def test_the_passport_names_row_and_column_as_numbers(tmp_path: Path) -> None:
    # M5: spec §6.7 lists `row, column`; the sheet had one `cell` like "C2".
    run = _run(tmp_path)
    assert main(["model", "--run", str(run)]) == 0
    passport = load_workbook(BytesIO((run / "model.xlsx").read_bytes()))["Passport"]
    assert [c.value for c in passport[1]][:4] == ["sheet", "row", "column", "line_id"]
    first = [c.value for c in passport[2]]
    assert isinstance(first[1], int) and isinstance(first[2], int)


def test_appending_text_rows_does_not_scan_existing_worksheet_cells(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from openpyxl import Workbook
    from openpyxl.worksheet.worksheet import Worksheet

    from rbotlib.model.xlsx import _row

    scans = []
    original = cast(property, Worksheet.max_row)
    assert original.fget is not None
    get_row = original.fget

    def counted(sheet: Worksheet) -> int:
        scans.append(sheet.title)
        return int(get_row(sheet))

    monkeypatch.setattr(Worksheet, "max_row", property(counted))
    sheet = Workbook().active
    assert isinstance(sheet, Worksheet)
    for n in range(100):
        _row(sheet, [f"={n}", n, None])
    assert scans == []
    assert sheet["A100"].value == "=99" and sheet["A100"].data_type == "s"
    assert sheet["B100"].value == 99
