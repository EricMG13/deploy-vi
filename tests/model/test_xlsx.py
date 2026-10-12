"""`model.xlsx` written from model.json's data and the tables' exact figures
(spec §6.7): values with number formats, text always as text."""

from __future__ import annotations

import zipfile
from collections.abc import Mapping, Sequence
from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from rbotlib.model.snapshot import Fig
from rbotlib.model.xlsx import FORMATS, SHEET_BANDS, workbook

COLUMN = dict.fromkeys("column_id kind label period_id fiscal_year end_date days case".split(), "x")
SOURCE = {"source": "10-K.txt", "locator": "p.61", "excerpt": "Total revenue 11,519"}


def _cell(value: str | None, reason: str | None = None) -> dict[str, Any]:
    return {
        "column_id": "FY:FY2025",
        "value": value,
        "basis": "REPORTED" if value else None,
        "reason": reason,
        "sources": [dict(SOURCE)] if value else [],
    }


def _model(label: str = "Total revenue", value: str | None = "11519.5") -> dict[str, Any]:
    line = {
        "line_id": "revenue",
        "band": 2,
        "label": label,
        "unit": "MONEY",
        "cells": [_cell(value)],
    }
    return {"columns": [{**COLUMN, "label": "FY25"}], "lines": [line]}


def _book(
    model: dict[str, Any],
    house: Sequence[Mapping[str, Any]] = (),
    snapshot: Sequence[Mapping[str, Any]] = (),
) -> Workbook:
    return load_workbook(BytesIO(workbook(model, house, snapshot)))


def test_the_sheets_come_in_the_specified_order() -> None:
    book = _book(_model())
    assert book.sheetnames == [
        "Columns", *SHEET_BANDS, "House", "Snapshot", "Projection", "Passport"
    ]  # fmt: skip


def test_a_band_sheet_holds_the_value_with_its_number_format_and_a_passport_row() -> None:
    book = _book(_model())
    sheet = book["02 Revenue by business unit"]
    assert [c.value for c in sheet[1]] == ["Line", "line_id", "FY25"]
    cell = sheet.cell(2, 3)
    assert (cell.value, cell.number_format) == (Decimal("11519.5"), FORMATS["MONEY"])
    passport = [[c.value for c in r] for r in book["Passport"].iter_rows(min_row=2)]
    assert passport == [
        [
            "02 Revenue by business unit",
            2,
            3,
            "revenue",
            "FY:FY2025",
            "REPORTED",
            None,
            "10-K.txt",
            "p.61",
            "Total revenue 11,519",
        ]
    ]


@pytest.mark.parametrize("lead", ["=", "+", "-", "@"])
def test_text_is_stored_as_a_string_whatever_its_first_character(lead: str) -> None:
    # Review I4: openpyxl stores any string starting with `=` as a formula.
    text = f'{lead}HYPERLINK("http://x.example","click")'
    model = _model(label=text)
    model["lines"][0]["cells"][0]["sources"][0]["excerpt"] = f"{lead}cmd|' /C calc'!A0"
    house = [{"item": text, "value": Fig("1", Decimal(1), "MONEY", None), "components": text}]
    data = workbook(model, house, [])
    with zipfile.ZipFile(BytesIO(data)) as archive:
        sheets = [n for n in archive.namelist() if n.startswith("xl/worksheets/sheet")]
        assert all("<f>" not in archive.read(n).decode() for n in sheets)
    book = load_workbook(BytesIO(data))
    assert book["02 Revenue by business unit"].cell(2, 1).value == text
    assert book["House"].cell(2, 1).value == text and book["House"].cell(2, 6).value == text
    assert all(c.data_type != "f" for s in book for r in s.iter_rows() for c in r)
    assert str(book["Passport"].cell(2, 10).value).startswith(f"{lead}cmd")


def test_house_and_snapshot_hold_the_exact_figures_and_a_gap_is_blank_with_its_reason() -> None:
    # Review I6: the sheets held display text ("4134", "98.43x").
    exact = Decimal("3.5920859794")
    house = [
        {"item": "SSIM Multiple", "value": Fig("3.59x", exact, "MULTIPLE", None), "reason": "—"},
        {"item": "EV", "value": Fig("—", None, "MONEY", "COMPONENT_MISSING"), "reason": "x"},
    ]
    snapshot = [
        {"line_id": "gross_profit_growth", "fy": Fig("—", None, "PERCENT", "NOT_APPLICABLE")}
    ]
    book = _book(_model(), house, snapshot)
    ssim = book["House"].cell(2, 2)
    assert (
        ssim.number_format == FORMATS["MULTIPLE"] and abs(Decimal(str(ssim.value)) - exact) < 1e-9
    )
    assert book["House"].cell(3, 2).value is None
    passport = [[c.value for c in r][:7] for r in book["Passport"].iter_rows(min_row=2)]
    assert ["House", 3, 2, "EV", "value", None, "COMPONENT_MISSING"] in passport
    assert ["Snapshot", 2, 7, "gross_profit_growth", "fy", None, "NOT_APPLICABLE"] in passport
