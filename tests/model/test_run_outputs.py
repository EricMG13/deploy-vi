"""`rbot model --run`: MODEL.md and model.xlsx (spec §3.2, §6.7)."""

from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from credit_model_fixture import example_registers
from openpyxl import load_workbook
from run_cases import _model, _run, cp1d_handoff

from rbotlib.cli import main
from rbotlib.handoff import markers, parse
from rbotlib.model.assemble import ModelInputs, credit_model
from rbotlib.model.output import Limitation, qa_status
from rbotlib.model.run import _model_limitations
from rbotlib.model.xlsx import SHEET_BANDS
from rbotlib.rules import load_rules


def test_model_md_is_a_handoff_whose_citations_resolve(tmp_path: Path) -> None:
    run = _run(tmp_path, cp1d=cp1d_handoff())
    main(["model", "--run", str(run)])
    text = (run / "handoffs" / "MODEL.md").read_text(encoding="utf-8")
    h = parse(text)
    rules = load_rules()
    assert set(rules.front_matter) <= set(h.front_matter)
    assert h.front_matter["module_id"] == "MODEL"
    assert h.front_matter["inputs"] == "CP-1, CP-1D, RUN_BRIEF"
    assert tuple(h.sections) == rules.sections
    assert h.duplicate_tags == [] and h.orphan_tags == []
    assert list(h.tables) == [
        "snapshot.capital_structure",
        "snapshot.capital_structure_totals",
        "snapshot.ebitda_adjustments",
        "snapshot.financial_columns",
        "snapshot.financials",
        "model.summary",
        "model.metrics",
        "model.house",
        "model.maturity_wall",
        "model.limitations",
        "citations",
    ]
    ids = {row[0] for row in h.tables["citations"].rows}
    assert markers(h.body) == ids
    assert ids == {f"C{n}" for n in range(1, len(ids) + 1)}
    house = {row[0]: row for row in h.tables["model.house"].rows}
    assert house["SSIM EBITDA"][1] == str(4094 - 32)  # CP-1D rejects restructuring
    assert house["Adjusted debt"][1] == str(14706 + 120)
    codes = [row[0] for row in h.tables["model.limitations"].rows]
    assert "CP1D_NOT_READ" not in codes and "NO_PROJECTIONS" in codes
    data = _model(run)
    lines = {line["line_id"]: line for line in data["lines"]}
    pension = {c["column_id"]: c for c in lines["debt_like:pension_deficit"]["cells"]}
    assert pension["PF:LTM_Q2_2026"]["sources"][0]["module"] == "CP-1D"
    assert pension["PF:LTM_Q2_2026"]["sources"][0]["excerpt"] == "pension deficit of $120 million"


def test_a_bridge_that_does_not_close_restricts_the_model() -> None:
    model = credit_model(ModelInputs(example_registers(), exclude_unrealised=True))
    index = [c.column_id for c in model.columns].index("FY:FY2024")
    lines = tuple(
        replace(
            x,
            cells=tuple(
                replace(c, value=Decimal(1)) if n == index else c for n, c in enumerate(x.cells)
            ),
        )
        if x.line_id == "ebit"
        else x
        for x in model.lines
    )
    found = _model_limitations(replace(model, lines=lines))
    fault = next(x for x in found if x.code == "RECONCILIATION_FAULT")
    assert fault == Limitation(
        "RECONCILIATION_FAULT", "FY:FY2024:ebitda does not close", "ebitda", "FY:FY2024"
    )
    assert qa_status(found) == "Restricted"
    assert qa_status([Limitation("NO_LTM", "x")]) == "Restricted"
    assert qa_status([Limitation("PF_LTM_EQUALS_LTM", "x")]) == "Pass"


def test_model_xlsx_holds_the_values_by_band(tmp_path: Path) -> None:
    run = _run(tmp_path)
    main(["model", "--run", str(run)])
    book = load_workbook(BytesIO((run / "model.xlsx").read_bytes()))
    assert book.sheetnames == [
        "Columns",
        *SHEET_BANDS,
        "House",
        "Snapshot",
        "Projection",
        "Passport",
    ]
    revenue_band = book["02 Revenue by business unit"]
    header = [c.value for c in revenue_band[1]]
    assert header[:2] == ["Line", "line_id"] and "FY25" in header
    row = next(r for r in revenue_band.iter_rows(min_row=2) if r[1].value == "revenue")
    cell = row[header.index("FY25")]
    assert (cell.value, cell.number_format) == (11519, "#,##0.0")
    metrics = book["12 Credit metrics"]
    multiple = next(
        r for r in metrics.iter_rows(min_row=2) if r[1].value == "metric:total_leverage_adjusted"
    )
    assert multiple[header.index("FY25")].number_format == '0.00"x"'
    passport = book["Passport"]
    assert [c.value for c in passport[1]] == [
        "sheet",
        "row",
        "column",
        "line_id",
        "column_id",
        "basis",
        "reason",
        "source",
        "locator",
        "excerpt",
    ]
    assert passport.max_row > 100
    assert all(
        not isinstance(c.value, str) or not c.value.startswith("=")
        for r in revenue_band.iter_rows()
        for c in r
    )
